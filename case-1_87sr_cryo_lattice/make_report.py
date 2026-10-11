from docx import Document
doc=Document()
doc.add_heading('87Sr 低温光晶格钟 — 三维装置渲染验证记录',0)
doc.add_paragraph('文件版本：v4 ｜ 日期：2026-10-10 ｜ 可编辑草稿')
doc.add_heading('一、本轮输入',1)
for t in ['用户反馈：v2 整图全灰、未通过；要求撤回"全部可见"自检结论并如实记录',
          '用户指定交付：无文字整体预览 + 铜腔局部剖切预览，.blend + 可重跑脚本，简短记录',
          '范围限定：Fig.7 的 MOT 与拟建低温铜腔局部，暂不扩展炉子/慢速器/移动平台']:
    doc.add_paragraph(t,style='List Bullet')
doc.add_heading('二、实际修改（build_blender_v4.py，基于 v3）',1)
for m in ['真空室：布尔薄壁壳体(3mm壁厚)，剖切前壁+顶部前1/4，端口开孔+玻璃窗',
 '低温铜腔：布尔真正空心（内径10mm=半径5mm，长160mm，论文值）；壁厚1.5mm为假设值；CuO内衬；aperture隔膜+3mm半径通孔(假设值)；后倾斜蓝宝石窗',
 'MOT线圈：两个对置环形 torus（外径64mm，线径5mm）',
 '冷却光束(461nm)×4；晶格光束(magic wavelength [11])双向；钟光束(698nm)+二向色镜(拟建)',
 '相机 cam1 总览 / cam3 侧视剖切；无 3D 文字']:
    doc.add_paragraph(m,style='List Bullet')
doc.add_heading('三、自动检查与覆盖范围（如实）',1)
rows=[['检查项','结果','覆盖范围限制'],
 ['相机视锥归一化坐标','cam1/cam3 下全部关键对象 in_full_frame=True','仅验证"在画面范围内"，不验证是否被前层实体遮挡'],
 ['aperture 通孔射线','射线穿入通孔后命中后方光束而非实心堵死','确认通孔为空心'],
 ['PNG 色彩丰富度','coarse colors 72–78，bright frac 0.23–0.41，非全灰','只排除整图单色故障，不能证明具体部件可见'],
 ['遮挡射线 cam→目标中心','cam_overview: mot_coil_L/R、ap_ring、sapph_window 被 cav_outer/lat_beam_1 先命中；cam_cut 类似','这是直线投影"被何物先命中"信息，不等于渲染图中看不见；脚本保留在 occlusion_check_v4.py，不据此声称无遮挡'],
 ['能否直接查看 PNG 像素','read 工具无法读二进制 PNG 内容','以上为几何/统计层面验证，最终可见性需用户目检']]
t=doc.add_table(rows=0,cols=3)
for r in rows:
    cells=t.add_row().cells
    for i,x in enumerate(r): cells[i].text=x
doc.add_heading('四、已知近似/未确定项',1)
for t2 in ['真空室外形 200×150×150mm、端口 20mm、MOT 线圈间距 40mm：论文未给，示意值',
 '腔壁厚 1.5mm、aperture 通孔半径 3mm：假设值，已记录',
 'magic wavelength 具体数值：未核实，图中不标具体 nm（[11]=Katori PRL 91:173005 2003）',
 '冷却温度序列：3 mK（一级MOT）→ 50 µK（二级MOT）→ 4 µK（装格前），正文核实',
 '装置状态：低温铜腔/后窗/dichroic = proposed（拟建），本次为三维重构，非已建成整体装置']:
    doc.add_paragraph(t2,style='List Bullet')
doc.add_heading('五、本轮执行结论（未获用户验收）',1)
doc.add_paragraph('自动检查通过"视锥入框"+"非全灰"两项；遮挡检查给出投影层面信息，不能替代目检。'
 '尚未获用户确认两张 PNG 实际可见内容；按"不能因模型含某部件就宣称图片中看得见它"的要求，'
 '记录到此，待用户目检反馈。')
doc.save('v4_review_report.docx')
print('OK')
